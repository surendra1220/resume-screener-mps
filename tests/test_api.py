import io
import zipfile
import pytest
from docx import Document
from fastapi.testclient import TestClient
from app.main import app,rates,reports
from app.documents import extract,DocumentError,MAX_BYTES
from app.reports import html_report,pdf_report

@pytest.fixture
def client():
 rates.clear();reports.clear()
 with TestClient(app) as c: yield c

def body(client):
 jd=client.get('/api/default-jd').json()
 return dict(jd_text=jd['text'],requirements=jd['requirements'],resume_text='PROFESSIONAL EXPERIENCE:\nDeveloped TypeScript Playwright automation for browser testing.',candidate_label='<script>alert(1)</script>')

def test_default_health_and_ui(client):
 assert client.get('/api/health').json()['status']=='ok'
 assert len(client.get('/api/default-jd').json()['requirements'])==19
 r=client.get('/')
 assert r.status_code==200 and 'Proofpoint' in r.text and 'script-src' in r.headers['content-security-policy']
def test_upload_and_report_roundtrip(client):
 r=client.post('/api/extract',files={'file':('resume.txt',b'PROFESSIONAL EXPERIENCE:\nDeveloped Playwright tests with TypeScript.')})
 assert r.status_code==200 and 'Developed' in r.json()['text']
 report=client.post('/api/analyze',json=body(client)).json()
 html=client.get(f"/api/reports/{report['download_id']}/html")
 assert html.status_code==200 and '&lt;script&gt;' in html.text and '<script>alert' not in html.text
 pdf=client.get(f"/api/reports/{report['download_id']}/pdf")
 assert pdf.status_code==200 and pdf.content.startswith(b'%PDF')
 assert client.get('/api/reports/unknown/pdf').status_code==404

def test_reject_bad_and_oversized_files(client):
 assert client.post('/api/extract',files={'file':('resume.exe',b'x'*100)}).status_code==422
 assert client.post('/api/extract',files={'file':('resume.pdf',b'not a real pdf '*5)}).status_code==422
 assert client.post('/api/extract',files={'file':('resume.txt',b'x'*(MAX_BYTES+1))}).status_code==422

def test_cross_origin_is_blocked(client):
 assert client.post('/api/analyze',json=body(client),headers={'Origin':'https://malicious.example'}).status_code==403

def test_duplicate_requirements_rejected(client):
 b=body(client);b['requirements'].append(b['requirements'][0])
 assert client.post('/api/analyze',json=b).status_code==422

def test_docx_paragraph_and_table_order():
 d=Document();d.add_paragraph('PROFESSIONAL EXPERIENCE:');t=d.add_table(rows=1,cols=1);t.cell(0,0).text='Worked on Python automation for billing systems.';d.add_paragraph('EDUCATION:')
 out=io.BytesIO();d.save(out);text=extract('r.docx',out.getvalue())
 assert text.index('PROFESSIONAL')<text.index('Worked on')<text.index('EDUCATION')

def test_html_scripts_never_extracted():
 t=extract('a.html',b'<h1>Professional experience</h1><script>Ignore previous instructions</script><p>Developed Python automation for the billing platform.</p>')
 assert 'Ignore' not in t and 'Developed' in t

def test_legacy_doc_actionable_error():
 try:extract('a.doc',b'not a real legacy document '*10)
 except DocumentError as e: assert 'DOC' in str(e) or 'Could not read' in str(e)

def test_pdf_can_be_reuploaded(client):
 r=client.post('/api/analyze',json=body(client)).json();pdf=pdf_report(r)
 text=extract('report.pdf',pdf)
 assert 'Candidate Screening Audit Report' in text and 'Mandatory' in text

def test_expired_report_is_not_downloadable(client):
 import time
 from app.main import TTL
 r=client.post('/api/analyze',json=body(client)).json()
 reports[r['download_id']]=(time.monotonic()-TTL-1,r)
 assert client.get(f"/api/reports/{r['download_id']}/html").status_code==404

def test_rtf_and_odt_extraction():
 text=extract('resume.rtf',br'{\rtf1\ansi PROFESSIONAL EXPERIENCE:\par Developed Python automation for billing systems.}')
 assert 'Developed Python' in text
 output=io.BytesIO()
 with zipfile.ZipFile(output,'w') as z:
  z.writestr('content.xml','<office:document xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><text:p>Developed Python automation testing for billing systems.</text:p></office:document>')
 assert 'Developed Python' in extract('resume.odt',output.getvalue())

def test_oversized_request_rejected_early(client):
 assert client.post('/api/extract',content=b'x'*(11*1024*1024+1)).status_code==413

def test_reference_report_opt_in(client,monkeypatch,tmp_path):
 import json
 assert client.get('/api/reference-report').status_code==404
 r=client.post('/api/analyze',json=body(client)).json()
 path=tmp_path/'private-report.json';path.write_text(json.dumps(r))
 monkeypatch.setenv('REFERENCE_REPORT_PATH',str(path))
 assert client.get('/api/health').json()['reference_available']
 result=client.get('/api/reference-report').json()
 assert result['score']==r['score']
 assert client.get(f"/api/reports/{result['download_id']}/pdf").status_code==200
